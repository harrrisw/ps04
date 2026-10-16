"""Create plain Word article documents from the delivered Markdown sources."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
FOLDER = ROOT/'deliverables/20261010T052912Z-1d2e8a'
TOKEN = re.compile(r'\[([^\]]+)\]\((https?://[^)]+)\)|\*\*(.+?)\*\*|\*([^*]+)\*')


def add_text(paragraph, text, bold=False, italic=False):
    position = 0
    for match in TOKEN.finditer(text):
        run = paragraph.add_run(text[position:match.start()])
        run.bold, run.italic = bold, italic
        label, url, strong, emphasis = match.groups()
        if url:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            run_xml = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'),'175D73'); props.append(color)
            underline = OxmlElement('w:u'); underline.set(qn('w:val'),'single'); props.append(underline)
            if italic:
                props.append(OxmlElement('w:i'))
            if bold:
                props.append(OxmlElement('w:b'))
            run_xml.append(props)
            words = OxmlElement('w:t'); words.text = label; run_xml.append(words)
            link.append(run_xml); paragraph._p.append(link)
        elif strong:
            add_text(paragraph,strong,True,italic)
        else:
            add_text(paragraph,emphasis,bold,True)
        position = match.end()
    run = paragraph.add_run(text[position:])
    run.bold, run.italic = bold, italic


for pid in ['01','02','03']:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.75)
    section.left_margin = section.right_margin = Inches(.85)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(11)
    normal.font.color.rgb = RGBColor(0,0,0)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.widow_control = True
    for name,size in [('Title',21),('Heading 1',14),('Heading 2',12)]:
        style = document.styles[name]
        style.font.name,style.font.size = 'Calibri',Pt(size)
        style.font.color.rgb = RGBColor(0,0,0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name!='Title' else 0)
        style.paragraph_format.space_after = Pt(7)
    source = (FOLDER/f'{pid}-article.md').read_text(encoding='utf-8')
    for block in re.split(r'\n\s*\n',source.strip()):
        heading = re.fullmatch(r'(#{1,3}) (.+)',block)
        if heading:
            style = 'Title' if len(heading[1])==1 else f'Heading {len(heading[1])-1}'
            p = document.add_paragraph(style=style)
            add_text(p,heading[2])
        elif all(line.startswith('- ') for line in block.splitlines()):
            for line in block.splitlines():
                p=document.add_paragraph(style='List Bullet');add_text(p,line[2:])
        else:
            p=document.add_paragraph()
            lines=block.splitlines()
            for n,line in enumerate(lines):
                add_text(p,line.rstrip())
                if n<len(lines)-1:
                    p.add_run().add_break() if line.endswith('  ') else p.add_run(' ')
    document.core_properties.title = source.splitlines()[0].removeprefix('# ')
    document.core_properties.author = 'The Small Business Research Brief'
    destination = FOLDER/f'{pid}-article.docx'
    document.save(destination)
    print(destination)
